"""Routes for the ITC STEM spreadsheet, the student's profile and the per-job "apply" page.

Thin like app.py: read the form, call the core, render a template. Everything that calls
the AI here (fit for one job, cover letter) runs in the request: it takes 10-40 s, so the
pages say so and the buttons show a waiting state (app.js).
"""

import re
from datetime import datetime
from urllib.parse import quote
from pathlib import Path

from fastapi import FastAPI, HTTPException, Request
from fastapi.responses import FileResponse, HTMLResponse, PlainTextResponse, RedirectResponse
from openpyxl import load_workbook
from starlette.datastructures import UploadFile

from cityu_jobs import config
from cityu_jobs.core.ai.providers import ProviderError
from cityu_jobs.core.apply_info import apply_info, clean_label, itc_page_url
from cityu_jobs.core.cache import Cache
from cityu_jobs.core.dates import job_timeline
from cityu_jobs.core.fit import assess_fit, can_assess
from cityu_jobs.core.itc_stem_scraper import fix_mojibake, manual_files, parse_job_list
from cityu_jobs.core.letters import STATUSES, LetterError, load_letter, load_tracker, set_status, write_letter
from cityu_jobs.core.profile import CVError, Profile, extract_cv_text
from cityu_jobs.core.settings import Settings

PROFILE_FIELDS = ("name", "email", "phone", "education", "skills", "projects", "experience", "extra", "cv_text")
_SAFE_NAME = re.compile(r"[^\w .()-]+")


async def _read_upload(upload, limit: int) -> bytes:
    """The uploaded file's bytes, refusing anything over `limit` without reading it all."""
    data = await upload.read(limit + 1)
    if len(data) > limit:
        raise HTTPException(status_code=413, detail=f"File is larger than {limit // (1024 * 1024)} MB.")
    return data


def _itc_files() -> list[Path]:
    """The ITC spreadsheets the app can read, newest (the one in use) first."""
    files = {f for s in config.ITC_STEM_SOURCES if s.get("enabled") for f in manual_files(s["college"])}
    return sorted(files, key=lambda f: f.stat().st_mtime, reverse=True)


def _itc_file(name: str) -> Path:
    """A file in .data/itc/ by name; never a path outside it."""
    for f in _itc_files():
        if f.name == name:
            return f
    raise HTTPException(status_code=404, detail="No such ITC STEM file.")


def read_sheet(path: Path, max_rows: int = 500) -> tuple[list[str], list[list[str]], str]:
    """(headers, rows, sheet name) of the sheet with the job list, as text cells."""
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        best: tuple[list[str], list[list[str]], str] = ([], [], "")
        for sheet in workbook.worksheets:
            rows = [["" if v is None else (v.date().isoformat() if isinstance(v, datetime) else fix_mojibake(str(v))) for v in row]
                    for row in sheet.iter_rows(values_only=True)]
            rows = [r for r in rows if any(c.strip() for c in r)]
            if not rows:
                continue
            # The header is the first row with several filled cells (lists may start with a title row).
            header_at = next((i for i, r in enumerate(rows[:15]) if sum(bool(c.strip()) for c in r) >= 3), 0)
            headers, body = rows[header_at], rows[header_at + 1 :][:max_rows]
            if len(body) > len(best[1]):
                width = max(len(headers), *(len(r) for r in body)) if body else len(headers)
                keep = [i for i in range(width) if (i < len(headers) and headers[i].strip()) or any(i < len(r) and r[i].strip() for r in body)]
                best = (
                    [headers[i] if i < len(headers) else "" for i in keep],
                    [[r[i] if i < len(r) else "" for i in keep] for r in body],
                    sheet.title,
                )
        return best
    finally:
        workbook.close()


def _back(path: str, error: str, anchor: str = "") -> RedirectResponse:
    """Redirect to `path` with an error message the page shows (URL-encoded)."""
    return RedirectResponse(f"{path}?error={quote(error[:200])}{anchor}", status_code=303)


def find_job(runner, job_id: str):
    """(job, ranked result or None, overall rank or None). Searches the last results, then the cache."""
    for rank, r in enumerate(runner.state.results, start=1):
        if r.job.job_id == job_id:
            return r.job, r, rank
    job = Cache.load().get_job(job_id)
    if job is None:
        raise HTTPException(status_code=404, detail="Job not found. It may have closed; run a new search.")
    return job, None, None


def register(app: FastAPI, render, runner, check_same_origin) -> None:
    # --- ITC STEM spreadsheet --------------------------------------------------------

    @app.get("/itc", response_class=HTMLResponse)
    def itc_page(request: Request, file: str = "", uploaded: int = -1, error: str = ""):
        files = _itc_files()
        current = _itc_file(file) if file else (files[0] if files else None)
        headers, rows, sheet = read_sheet(current) if current else ([], [], "")
        number_col = next((i for i, h in enumerate(headers) if h.strip().lower() in ("no.", "no", "ref", "#")), None)
        # Long-text columns (descriptions) get more width; short ones (No., dates) stay narrow.
        wide = [sum(len(r[i]) for r in rows) / max(len(rows), 1) > 60 for i in range(len(headers))]
        return render(
            request, "itc.html",
            files=[{"name": f.name, "size": f.stat().st_size, "mtime": datetime.fromtimestamp(f.stat().st_mtime)} for f in files],
            current=current.name if current else None,
            in_use=files[0].name if files else None,
            headers=headers, labels=[clean_label(h) for h in headers], wide=wide,
            rows=rows, sheet=sheet, number_col=number_col,
            itc_page_url=itc_page_url(), uploaded=uploaded, error=error,
        )

    @app.get("/itc/file/{name}")
    def itc_download(name: str):
        path = _itc_file(name)
        return FileResponse(path, filename=path.name, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

    @app.post("/itc/upload")
    async def itc_upload(request: Request):
        check_same_origin(request)
        form = await request.form()
        upload = form.get("file")
        if not isinstance(upload, UploadFile) or not upload.filename:
            return _back("/itc", "Choose an .xlsx file first.")
        data = await _read_upload(upload, config.ITC_MAX_BYTES)
        name = _SAFE_NAME.sub("_", Path(upload.filename).name).strip() or "job_list.xlsx"
        if not name.lower().endswith(".xlsx") or not data.startswith(b"PK"):
            return _back("/itc", "That is not an .xlsx spreadsheet.")
        try:
            jobs = parse_job_list(data, config.ITC_STEM_SOURCES[0])
        except Exception:  # openpyxl raises many kinds on broken files
            return _back("/itc", "The spreadsheet could not be read.")
        if not jobs:
            return _back("/itc", "No job rows were found in that spreadsheet.")
        config.ITC_DIR.mkdir(parents=True, exist_ok=True)
        (config.ITC_DIR / name).write_bytes(data)  # newest file wins; the next search re-reads it
        return RedirectResponse(f"/itc?uploaded={len(jobs)}", status_code=303)

    # --- Profile --------------------------------------------------------------------

    @app.get("/profile", response_class=HTMLResponse)
    def profile_page(request: Request, saved: bool = False, cv: int = 0, error: str = ""):
        profile = Profile.load()
        settings = Settings.load()
        return render(
            request, "profile.html", profile=profile, saved=saved, cv_loaded=cv, error=error,
            uses_claude=settings.ai_provider == "claude", extensions=", ".join(config.CV_EXTENSIONS),
        )

    @app.post("/profile")
    async def profile_save(request: Request):
        check_same_origin(request)
        form = await request.form()
        profile = Profile.load()
        for name in PROFILE_FIELDS:
            value = form.get(name)
            if isinstance(value, str):
                setattr(profile, name, value.strip()[: config.PROFILE_MAX_CHARS])
        profile.save()
        return RedirectResponse("/profile?saved=1", status_code=303)

    @app.post("/profile/cv")
    async def profile_cv(request: Request):
        check_same_origin(request)
        form = await request.form()
        upload = form.get("cv")
        if not isinstance(upload, UploadFile) or not upload.filename:
            return _back("/profile", "Choose a CV file first.")
        data = await _read_upload(upload, config.CV_MAX_BYTES)
        try:
            text = extract_cv_text(upload.filename, data)
        except CVError as e:
            return _back("/profile", f"Could not read the CV: {e}.")
        profile = Profile.load()
        profile.cv_text, profile.cv_filename = text, Path(upload.filename).name[:120]
        profile.save()  # only the text is kept, never the uploaded file itself
        return RedirectResponse(f"/profile?cv={len(text)}", status_code=303)

    @app.post("/profile/clear")
    async def profile_clear(request: Request):
        check_same_origin(request)
        config.PROFILE_FILE.unlink(missing_ok=True)
        return RedirectResponse("/profile?saved=1", status_code=303)

    # --- One job: all details, how to apply, difficulty for you, cover letter ---------

    @app.get("/job/{job_id:path}/letter.txt")
    def letter_text(job_id: str):
        letter = load_letter(job_id)
        if not letter:
            raise HTTPException(status_code=404, detail="No letter written for this job yet.")
        return PlainTextResponse(
            f"{letter['letter']}\n",
            headers={"content-disposition": f'attachment; filename="cover-letter-{re.sub(r"[^A-Za-z0-9]+", "-", letter["title"])[:60]}.txt"'},
        )

    @app.post("/job/{job_id:path}/letter")
    async def letter_write(request: Request, job_id: str):
        check_same_origin(request)
        form = await request.form()
        job, ranked, _rank = find_job(runner, job_id)
        profile = Profile.load()
        if profile.is_empty():
            return _back(f"/job/{job_id}", "Add your CV or projects on the Profile page first.", "#letter")
        timeline = ranked.timeline if ranked and ranked.timeline else job_timeline(job)
        notes = str(form.get("notes") or "")[:1500]
        try:
            write_letter(job, apply_info(job, timeline.deadline), profile, Settings.load().provider(), notes)
        except (ProviderError, LetterError) as e:
            return _back(f"/job/{job_id}", str(e), "#letter")
        return RedirectResponse(f"/job/{job_id}?written=1#letter", status_code=303)

    @app.post("/job/{job_id:path}/fit")
    async def job_fit(request: Request, job_id: str):
        check_same_origin(request)
        job, ranked, _rank = find_job(runner, job_id)
        if runner.busy:  # the search also writes the cache
            return _back(f"/job/{job_id}", "A search is running; try again when it has finished.", "#fit")
        profile, provider = Profile.load(), Settings.load().provider()
        if profile.is_empty() or not can_assess(provider):
            return _back(f"/job/{job_id}", "This needs your profile (Profile page) and Claude (Settings page).", "#fit")
        cache = Cache.load()
        try:
            fits = assess_fit([job], cache, provider, profile)
        except ProviderError as e:
            return _back(f"/job/{job_id}", str(e), "#fit")
        cache.save()
        if ranked is not None:
            ranked.fit = fits.get(job.job_id)
        return RedirectResponse(f"/job/{job_id}#fit", status_code=303)

    @app.post("/job/{job_id:path}/status")
    async def job_status(request: Request, job_id: str):
        check_same_origin(request)
        form = await request.form()
        job, _ranked, _rank = find_job(runner, job_id)
        status = str(form.get("status") or "")
        if status not in STATUSES:
            raise HTTPException(status_code=400, detail="Unknown status.")
        set_status(job, status)
        back = str(form.get("back") or "")
        return RedirectResponse(back if back.startswith("/results") else f"/job/{job_id}", status_code=303)

    @app.get("/job/{job_id:path}", response_class=HTMLResponse)
    def job_page(request: Request, job_id: str, written: bool = False, error: str = ""):
        job, ranked, rank = find_job(runner, job_id)
        cache_level = None
        timeline = ranked.timeline if ranked and ranked.timeline else None
        fit = ranked.fit if ranked else None
        if timeline is None or fit is None:
            cache = Cache.load()
            cache_level = cache.get_level(job.job_id, job.content_hash) or {}
            timeline = timeline or job_timeline(job, cache_level.get("dates"))
            profile = Profile.load()
            if not profile.is_empty():
                fit = fit or cache.get_fit(job.job_id, job.content_hash, profile.fingerprint())
        level = ranked.level if ranked else (cache_level or {}).get("value")
        reason = ranked.level_reason if ranked else (cache_level or {}).get("reason", "")
        skills = ranked.skills if ranked else (cache_level or {}).get("skills", [])
        profile = Profile.load()
        return render(
            request, "job.html",
            job=job, rank=rank, info=apply_info(job, timeline.deadline if timeline else None), timeline=timeline,
            fit=fit, level=level, level_reason=reason, skills=skills,
            letter=load_letter(job.job_id), written=written, error=error,
            status=(load_tracker().get(job.job_id) or {}).get("status", ""),
            profile_ready=not profile.is_empty(), can_fit=can_assess(Settings.load().provider()),
            busy=runner.busy, itc_number=(job.raw or {}).get("No.") if job.source == "itc_stem" else None,
        )

