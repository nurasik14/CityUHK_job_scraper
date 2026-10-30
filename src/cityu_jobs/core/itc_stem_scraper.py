"""ITC STEM internship jobs: one Excel (.xlsx) job list per college.

Two ways to get a list:
1. Online: open the college's ITC STEM page, find the .xlsx link, download it with the same
   browser. (Found 2026-10-25: the College of Computing file redirects to a CityU portal
   login, so this usually fails with ItcDownloadBlocked. We never try to log in or get
   around it.)
2. Manual fallback: the student downloads the file in their own browser and saves it as
   `.data/itc/<college>_job_list.xlsx` (e.g. `cc_job_list.xlsx`).

Parsing (`parse_job_list`) is pure: bytes in, Jobs out, so it can be tested offline. The
column layout isn't known yet, so columns are matched by *header names* through a small
alias table instead of by position; columns we don't recognize still land in `Job.raw`.
"""

import hashlib
import io
import re
import time
from datetime import date, datetime
from pathlib import Path
from urllib.parse import urljoin

from openpyxl import load_workbook
from playwright.sync_api import BrowserContext, Page

from cityu_jobs import config
from cityu_jobs.core.cache import Cache
from cityu_jobs.core.models import Job, compute_content_hash

SOURCE = "itc_stem"
SCHEME = "ITC STEM Internship"
# Stored in every ITC job's raw, so part of its content_hash. Bump when parse_job_list() builds Jobs differently, so
# cached levels and embeddings made from the old text are redone once.
ITC_PARSER_VERSION = 3
# "About the company" text added to the description is cut to about this many characters.
OVERVIEW_MAX_CHARS = 500


class ItcDownloadBlocked(RuntimeError):
    """The spreadsheet could not be downloaded automatically (login wall, bot check, ...)."""


# --- Fetching -----------------------------------------------------------------

def find_job_list_url(page: Page) -> str | None:
    """Absolute URL of the first .xlsx link on a loaded college page, or None."""
    for link in page.locator("a[href]").all():
        href = link.get_attribute("href") or ""
        if href.lower().split("?")[0].endswith(".xlsx"):
            return urljoin(page.url, href)
    return None


def download_xlsx(context: BrowserContext, url: str) -> bytes:
    """Download with the browser's cookies. Raises ItcDownloadBlocked instead of following a login."""
    response = context.request.get(url, max_redirects=0)
    body = response.body()
    if 300 <= response.status < 400:
        where = response.headers.get("location", "")
        reason = "needs a CityU login" if "login" in where.lower() else f"redirects to {where}"
        raise ItcDownloadBlocked(f"download {reason}")
    if not response.ok:
        raise ItcDownloadBlocked(f"download failed with HTTP {response.status}")
    if not body.startswith(b"PK"):  # .xlsx files are zip archives, which start with "PK"
        raise ItcDownloadBlocked("download returned a web page, not a spreadsheet (bot check?)")
    return body


def manual_files(college: str) -> list[Path]:
    """Spreadsheets the student saved for this college in .data/itc/, newest first.

    `<college>_*.xlsx` always counts. For the primary source (the first in
    config.ITC_STEM_SOURCES), any other .xlsx there counts too, so a file saved under
    its original name (e.g. "2026.10.22 Job List.xlsx") still works.
    """
    if not config.ITC_DIR.is_dir():
        return []
    prefixes = [f"{s['college']}_" for s in config.ITC_STEM_SOURCES]
    is_primary = college == config.ITC_STEM_SOURCES[0]["college"]
    files = [
        f
        for f in config.ITC_DIR.glob("*.xlsx")
        if not f.name.startswith("~$")  # Excel's lock files
        and (f.name.startswith(f"{college}_") or (is_primary and not f.name.startswith(tuple(prefixes))))
    ]
    return sorted(files, key=lambda f: f.stat().st_mtime, reverse=True)


# --- Parsing ------------------------------------------------------------------

# Field -> header keywords. Each keyword must match whole words in the lower-cased header
# (an "s" plural is allowed), so "ref" doesn't match "Preferred" and "post" doesn't match
# "Postgraduate". A trailing "*" marks a word stem: "responsibilit*" matches "Responsibilities".
# Fields are tried in this order and each column goes to the first field that matches, so
# specific fields ("closing date", "no. of positions") come before generic ones ("position").
_FIELD_ALIASES: list[tuple[str, list[str]]] = [
    ("honorarium", ["honorarium"]),
    ("location", ["location", "work place", "workplace", "venue"]),
    ("ref", ["ref", "reference", "job no", "job code", "job id", "post no", "position no"]),
    ("closing_date", ["closing", "deadline", "apply by", "application due"]),
    ("weekly_hours", ["per week", "weekly", "hour*", "hrs"]),
    ("positions", ["no. of", "number of", "vacanc*", "quota", "place", "headcount"]),
    ("part_time", ["part-time", "full-time", "part time", "full time", "employment type", "mode"]),
    ("duration", ["duration"]),
    ("period", ["period", "start date", "internship date"]),
    ("title", ["job title", "position", "project title", "title", "post"]),
    ("overview", ["overview", "about the company", "company profile", "company background"]),
    ("company", ["company", "organization", "organisation", "employer"]),
    ("department", ["department", "dept", "unit"]),
    ("nature", ["nature", "category", "field", "area"]),
    ("requirements", ["requirement", "qualification", "skill", "discipline", "major", "programme"]),
    ("description", ["description", "duties", "responsibilit*", "content", "detail", "scope"]),
]


def _keyword_pattern(keyword: str) -> re.Pattern:
    if keyword.endswith("*"):
        return re.compile(rf"(?<!\w){re.escape(keyword[:-1])}")
    return re.compile(rf"(?<!\w){re.escape(keyword)}s?(?!\w)")


# Headers that only mean something on their own: a bare "Date" column is the posting
# date (the CC list), but "date" inside a header ("Start Date") must not match it.
_EXACT_HEADERS = {"date": "posted", "posting date": "posted", "date posted": "posted"}

_FIELD_PATTERNS = [(name, [_keyword_pattern(k) for k in keywords]) for name, keywords in _FIELD_ALIASES]
# How many top rows to search for the header row (lists often start with a banner title).
_HEADER_SEARCH_ROWS = 15


# Text that was UTF-8 but got decoded as Windows-1252 somewhere upstream ("€“" for "–").
_MOJIBAKE = (("â€“", "–"), ("â€”", "—"), ("â€™", "’"), ("â€œ", "“"), ("â€\x9d", "”"), ("â€˜", "‘"),
             ("€“", "–"), ("€”", "—"), ("\u00c2\u00a0", " "))


def fix_mojibake(text: str) -> str:
    for bad, good in _MOJIBAKE:
        text = text.replace(bad, good)
    return text


def _cell_text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat() if value.time() == datetime.min.time() else value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    # "Â" + no-break space is a no-break space decoded with the wrong encoding upstream.
    return " ".join(fix_mojibake(str(value)).split())


# Free-text fields that may be spread over several columns (e.g. "Preferred Major" and
# "Skills Required"); all their columns are joined. Every other field uses one column.
_MULTI_COLUMN_FIELDS = {"requirements", "description"}


def _map_columns(headers: list[str]) -> dict[str, list[int]]:
    """{field: [column indexes]} for the headers we recognize."""
    mapping: dict[str, list[int]] = {}
    for index, header in enumerate(headers):
        text = header.lower()
        exact = _EXACT_HEADERS.get(text.strip())
        if exact:
            mapping.setdefault(exact, [index])
            continue
        for field_name, patterns in _FIELD_PATTERNS:
            taken = field_name in mapping and field_name not in _MULTI_COLUMN_FIELDS
            if not taken and any(p.search(text) for p in patterns):
                mapping.setdefault(field_name, []).append(index)
                break
    return mapping


def _find_header_row(rows: list[tuple]) -> int | None:
    """Index of the first row that looks like a header: has a title column plus one more known column."""
    for i, row in enumerate(rows[:_HEADER_SEARCH_ROWS]):
        mapping = _map_columns([_cell_text(v) for v in row])
        if "title" in mapping and len(mapping) >= 2:
            return i
    return None


def _number(value) -> float | None:
    """20 -> 20.0; "20 hours/week" -> 20.0; "Full-time" -> None."""
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return float(value)
    match = re.search(r"\d+(\.\d+)?", _cell_text(value))
    return float(match.group()) if match else None


def _date(value) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    text = _cell_text(value)
    for fmt in ("%Y-%m-%d", "%d-%b-%Y", "%d %b %Y", "%d %B %Y", "%d/%m/%Y", "%Y.%m.%d"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    return None


def _unswap_date(value, posted: date | None):
    """Undo a day/month swap in an Excel date cell (seen in the CC list of 2026-10-22).

    Some rows were typed day-first ("6/11/2026") into cells that Excel read month-first,
    so they hold 11 June instead of 6 November, while the text dates in the same sheet
    are day-first. A date cell earlier than the row's posting date, whose day could be a
    month (<= 12), and that becomes on/after the posting date when swapped, is swapped.
    If it can't be swapped that way, it is wrong either way: None, so the job stays visible.
    Text cells and rows without a posting date are left alone.
    """
    if not isinstance(value, datetime) or posted is None or value.date() >= posted:
        return value
    if value.day <= 12:
        swapped = value.replace(month=value.day, day=value.month)
        if swapped.date() >= posted:
            return swapped
    return None


def _shorten(text: str, limit: int) -> str:
    """Cut `text` to at most `limit` characters at a word boundary, ending with "…" if cut."""
    if len(text) <= limit:
        return text
    cut = text[: limit - 1]
    if " " in cut:
        cut = cut[: cut.rindex(" ")]
    return cut.rstrip(" ,;:.-") + "…"


def _offers_honorarium(text: str) -> bool:
    """False for the many ways the CC list says "none": "NA", "No.", "0", "... will not provide ..."."""
    word = re.sub(r"[^a-z0-9/]", "", text.lower())
    if word in ("", "0", "no", "na", "n/a", "nil", "none"):
        return False
    return not re.search(r"^no\b|\bnot provide\b|\bwill not\b", text.lower())


def _split_department(text: str) -> tuple[str | None, str | None]:
    """ "CS - Computer Science" -> ("CS", "Computer Science"); "Computer Science" -> (None, "Computer Science")."""
    match = re.match(r"^([A-Z]{2,6})\s*(?:--?|–|:)\s*(.+)$", text)
    if match:
        return match.group(1), match.group(2).strip()
    if re.fullmatch(r"[A-Z]{2,6}", text):
        return text, None
    return None, text or None


def _job_suffix(ref: str, title: str, company: str, department: str) -> str:
    """The pluggable part of `itc:<college>:<suffix>`: the reference number if the
    list has one, else a short hash of title + company + department (stable across runs
    as long as those don't change)."""
    if ref:
        return re.sub(r"\s+", "", ref)
    key = "|".join([title, company, department]).lower()
    return hashlib.sha256(key.encode("utf-8")).hexdigest()[:12]


def parse_job_list(xlsx_bytes: bytes, source: dict, url: str | None = None) -> list[Job]:
    """Spreadsheet bytes -> Jobs. `source` is one entry of config.ITC_STEM_SOURCES."""
    workbook = load_workbook(io.BytesIO(xlsx_bytes), read_only=True, data_only=True)
    jobs: list[Job] = []
    seen_ids: set[str] = set()
    for sheet in workbook.worksheets:
        rows = list(sheet.iter_rows(values_only=True))
        header_index = _find_header_row(rows)
        if header_index is None:
            continue  # e.g. a notes or instructions sheet
        headers = [_cell_text(v) for v in rows[header_index]]
        columns = _map_columns(headers)

        for row in rows[header_index + 1 :]:
            def get(field_name: str):
                """The cell for a single-column field, or None."""
                i = columns.get(field_name, [None])[0]
                return row[i] if i is not None and i < len(row) else None

            def joined(field_name: str) -> str:
                """All columns of a multi-column field as "Header: value" lines (just the value if only one)."""
                indexes = [i for i in columns.get(field_name, []) if i < len(row) and _cell_text(row[i])]
                if len(indexes) == 1:
                    return _cell_text(row[indexes[0]])
                return "\n".join(f"{headers[i]}: {_cell_text(row[i])}" for i in indexes)

            title = _cell_text(get("title"))
            if not title:
                continue  # blank or spacer row
            raw = {h or f"column {i + 1}": _cell_text(v) for i, (h, v) in enumerate(zip(headers, row)) if _cell_text(v)}
            raw["_sheet"] = sheet.title
            raw["_parser_version"] = ITC_PARSER_VERSION  # in content_hash: a parser change redoes AI work

            company = _cell_text(get("company"))
            dept_code, dept_name = _split_department(_cell_text(get("department")))
            if dept_code is None and dept_name is None:
                dept_code, dept_name = source["department_code"], source["name"]
            mode = _cell_text(get("part_time")).lower()
            posted = _date(get("posted"))

            suffix = _job_suffix(_cell_text(get("ref")), title, company, _cell_text(get("department")))
            job_id = f"itc:{source['college']}:{suffix}"
            if job_id in seen_ids:  # two identical rows: keep both, distinguishable
                job_id = f"{job_id}-{len(jobs) + 1}"
            seen_ids.add(job_id)

            overview = _shorten(_cell_text(get("overview")), OVERVIEW_MAX_CHARS)
            description_parts = [
                f"Company: {company}" if company else "",
                joined("description"),
                f"About the company: {overview}" if overview else "",
            ]
            period_parts = [_cell_text(_unswap_date(get("period"), posted)), _cell_text(get("duration"))]
            honorarium = _cell_text(get("honorarium"))
            if _offers_honorarium(honorarium):
                raw["_honorarium"] = honorarium  # not scored; shown in the results' Details
            jobs.append(
                Job(
                    job_id=job_id,
                    source=SOURCE,
                    scheme=SCHEME,
                    title=title,
                    department_code=dept_code,
                    department_name=dept_name,
                    job_nature=_cell_text(get("nature")) or None,
                    description="\n\n".join(p for p in description_parts if p),
                    requirements=joined("requirements"),
                    part_time=("part" in mode) if ("part" in mode or "full" in mode) else None,
                    weekly_hours=_number(get("weekly_hours")),
                    total_hours=None,
                    employment_period=", ".join(p for p in period_parts if p) or None,
                    closing_date=_date(_unswap_date(get("closing_date"), posted)),
                    positions=int(n) if (n := _number(get("positions"))) is not None else None,
                    url=url or source["page_url"],
                    company=company or None,
                    location=_cell_text(get("location")) or None,
                    raw=raw,
                    content_hash=compute_content_hash(raw),
                )
            )
    workbook.close()
    return jobs


# --- Whole run ----------------------------------------------------------------

def fetch_source(context: BrowserContext | None, source: dict) -> tuple[list[Job], str]:
    """Jobs for one college plus a one-line note for the student on where they came from.

    Tries the online list first (if a browser context is given), then the manual file.
    Raises ItcDownloadBlocked if neither works, with instructions in the message.
    """
    problem = "no browser"
    if context is not None:
        page = context.new_page()
        try:
            page.goto(source["page_url"], wait_until="domcontentloaded")
            url = find_job_list_url(page)
            if url is None:
                problem = "no .xlsx link on the page"
            else:
                time.sleep(config.ITC_DELAY_S)
                jobs = parse_job_list(download_xlsx(context, url), source, url)
                return jobs, f"{source['name']}: downloaded {len(jobs)} jobs"
        except ItcDownloadBlocked as e:
            problem = str(e)
        except Exception as e:  # network trouble etc.: still try the manual file
            problem = f"{type(e).__name__}: {e}"
        finally:
            page.close()

    files = manual_files(source["college"])
    if files:
        jobs = parse_job_list(files[0].read_bytes(), source)
        return jobs, f"{source['name']}: {len(jobs)} jobs from {files[0].name} ({problem})"
    raise ItcDownloadBlocked(
        f"{source['name']}: {problem}. Download the job list from {source['page_url']} in your "
        f"browser and save it as {config.ITC_DIR / (source['college'] + '_job_list.xlsx')}"
    )


def scrape_itc_stem(context: BrowserContext | None, cache: Cache, sources: list[dict] | None = None) -> list[str]:
    """Fetch every enabled ITC STEM source into the cache; returns one note per source.

    All ITC colleges share the source key "itc_stem", so jobs are only marked inactive
    (via mark_scrape_complete) when *every* enabled college was fetched.
    """
    sources = [s for s in (sources or config.ITC_STEM_SOURCES) if s.get("enabled")]
    notes: list[str] = []
    seen: set[str] = set()
    complete = True
    for source in sources:
        try:
            jobs, note = fetch_source(context, source)
        except ItcDownloadBlocked as e:
            notes.append(str(e))
            complete = False
            continue
        for job in jobs:
            cache.upsert_job(job)
            seen.add(job.job_id)
        notes.append(note)
    if complete and sources:
        cache.mark_scrape_complete(SOURCE, seen)
    return notes
